"""Generación de plantillas oficiales de importación masiva de Amazon (.xlsm) para el proveedor Signes."""
import io
import re
from pathlib import Path
import openpyxl
from api._weights import weight_grams

TEMPLATE_CANDIDATES = [
    Path(__file__).parent / "data" / "Amazon Bulk Import Creation.xlsm",
    Path(__file__).parent.parent / "Amazon Bulk Import Creation.xlsm",
    Path("Amazon Bulk Import Creation.xlsm"),
]


def get_template_path():
    for p in TEMPLATE_CANDIDATES:
        if p.is_file():
            return p
    raise FileNotFoundError("No se encontró la plantilla 'Amazon Bulk Import Creation.xlsm'")


def build_sku(raw_ref):
    if not raw_ref:
        return ""
    ref = str(raw_ref).strip()
    if not ref:
        return ""
    if ref.upper().endswith("SGI"):
        return ref[:-3].rstrip() + "SGI"
    return f"{ref}SGI"


def parse_dimension_parts(medidas_str):
    """
    Extrae altura, grosor y ancho a partir del campo de medidas (ej: '25x12x12 cm', '45x34x10 cm', '60x40 cm').
    Retorna un diccionario con:
      - 'height': número (altura)
      - 'depth': número (grosor)
      - 'width': número (ancho)
      - 'unit': 'Centímetros'
    o None si no hay medidas válidas.
    """
    if not medidas_str or str(medidas_str).strip() in ("-", "", "None", "desconocido", "no disponible"):
        return None
    raw_nums = re.findall(r'(\d+(?:[.,]\d+)?)', str(medidas_str))
    if not raw_nums:
        return None
    nums = []
    for n in raw_nums:
        try:
            val = float(n.replace(',', '.'))
            if val > 0:
                nums.append(int(val) if val.is_integer() else round(val, 2))
        except ValueError:
            pass
    if not nums:
        return None

    unit = "Centímetros"

    if len(nums) >= 3:
        n1, n2, n3 = nums[0], nums[1], nums[2]
        # Si el último número es claramente la altura mayor y los dos primeros son iguales (ej: 12x12x25 cm)
        if n3 > n1 and n1 == n2:
            height, width, depth = n3, n1, n2
        # Formato habitual español: Alto x Ancho x Fondo/Grosor (ej: 25x12x12 cm)
        else:
            height, width, depth = n1, n2, n3
    elif len(nums) == 2:
        # Formato de 2 dimensiones (ej: 60x40 cm para alfombra, cuadro o mantel plano)
        height, width, depth = nums[0], nums[1], 1
    else:
        height, width, depth = nums[0], nums[0], nums[0]

    return {
        "height": height,
        "width": width,
        "depth": depth,
        "unit": unit
    }


def generate_bulk_import_excel(items):
    """Rellena la plantilla Amazon Bulk Import Creation.xlsm con los datos generados para el proveedor Signes.

    Columnas según especificación:
    - Col A (1): SKU = ref + 'SGI'
    - Col B (2): Tipo de producto = 'Home'
    - Col C (3): Acción de listing = 'Crear o reemplazar (actualización completa)'
    - Col G (7): Nombre del producto = título generado por la IA
    - Col H (8): Destacado del artículo = highlights generados por la IA
    - Col I (9): Marca = 'ROCKING GIFTS'
    - Col J (10): Tipo de identificador del producto = 'EAN'
    - Col L (12): Nodos recomendados de búsqueda = número de nodo Amazon
    """
    if not isinstance(items, list):
        raise ValueError("Los elementos a exportar deben ser una lista")

    template_path = get_template_path()
    wb = openpyxl.load_workbook(template_path, keep_vba=True)
    if "Plantilla" not in wb.sheetnames:
        raise ValueError("La hoja 'Plantilla' no existe en el archivo de importación masiva")

    # Asegurar que no exista la pestaña de exploración de nodos para evitar rechazos en Amazon
    for sname in list(wb.sheetnames):
        if any(k in sname.lower() for k in ["explorar datos", "exploración", "exploracion", "browse data"]):
            del wb[sname]

    ws = wb["Plantilla"]
    weight_columns = {key: next((c.column for c in ws[5] if str(c.value).startswith('item_weight[') and str(c.value).endswith('#1.' + key)), None) for key in ('value', 'unit')}

    # Conservar ejemplos; los metadatos de Amazon indican dataRow=7.
    if ws.max_row >= 7:
        ws.delete_rows(7, ws.max_row - 6)

    for idx, item in enumerate(items):
        if not isinstance(item, dict):
            continue
        r = 7 + idx
        sku = build_sku(item.get("ref") or item.get("sku") or "")
        title = str(item.get("title") or "")
        highlights = str(item.get("highlights") or "")
        node_id = str(item.get("node_id") or "")

        divisor = 1
        try:
            raw_div = item.get("divisor")
            if raw_div is not None and str(raw_div).strip():
                divisor = max(1, int(raw_div))
        except (ValueError, TypeError):
            divisor = 1

        brand = "ROCKING GIFTS"
        ws.cell(row=r, column=1, value=sku)
        ws.cell(row=r, column=2, value="Home")
        ws.cell(row=r, column=3, value="Crear o reemplazar (actualización completa)")
        ws.cell(row=r, column=7, value=title)
        ws.cell(row=r, column=8, value=highlights)
        ws.cell(row=r, column=9, value=brand)
        ws.cell(row=r, column=10, value="EAN")
        ws.cell(row=r, column=12, value=node_id)

        # Col Q (17): Nivel de paquete
        ws.cell(row=r, column=17, value="Unidad")

        # Col R (18): El paquete contiene la cantidad de SKU (número de unidades en el pack)
        ws.cell(row=r, column=18, value=divisor)

        # Col BA (53): Número de artículos
        ws.cell(row=r, column=53, value=divisor)

        # Col T (20): Número de modelo (mismo valor que Columna A)
        ws.cell(row=r, column=20, value=sku)

        # Col U (21): Fabricante (mismo valor que Columna I)
        ws.cell(row=r, column=21, value=brand)

        # Col X (24): URL de la imagen principal
        # Col AG (33): URL de la imagen de la muestra
        image_url = str(item.get("image_url") or item.get("url") or "").strip()
        if image_url:
            ws.cell(row=r, column=24, value=image_url)
            ws.cell(row=r, column=33, value=image_url)

        # Col AN (40): Descripción del producto generada por la IA
        description = str(item.get("description") or "")
        ws.cell(row=r, column=40, value=description)

        # Col AO (41), AP (42), AQ (43), AR (44), AS (45): Bullets 1 a 5
        bullets = item.get("bullet_points")
        if not isinstance(bullets, list):
            bullets = [item.get(f"bullet_{i+1}", "") for i in range(5)]
        for b_idx, col_num in enumerate([41, 42, 43, 44, 45]):
            b_val = str(bullets[b_idx]) if b_idx < len(bullets) and bullets[b_idx] else ""
            ws.cell(row=r, column=col_num, value=b_val)

        # Col AT (46): Palabra clave genérica / Backend Keywords (Search Terms, máx 250 car.)
        keywords = str(item.get("backend_keywords") or "").strip()
        keywords = re.sub(r'[,;.:_"\'-]+', ' ', keywords)
        keywords = re.sub(r'\s+', ' ', keywords).strip()
        if len(keywords) > 250:
            truncated = keywords[:250]
            last_space = truncated.rfind(' ')
            keywords = truncated[:last_space].strip() if last_space > 150 else truncated.strip()
        ws.cell(row=r, column=46, value=keywords)

        # Col AV (48): Material
        material = str(item.get("material") or "").strip()
        if material.lower() in ("-", "desconocido", "no disponible", "n/a", "none", "null"):
            material = ""
        ws.cell(row=r, column=48, value=material)

        # Col BD (56): Color
        color = str(item.get("color") or "").strip()
        if color.lower() in ("-", "desconocido", "no disponible", "n/a", "none", "null"):
            color = ""
        ws.cell(row=r, column=56, value=color)

        # Col BE (57): Medidas / Tamaño
        medidas = str(item.get("medidas") or "").strip()
        if medidas.lower() in ("-", "desconocido", "no disponible", "n/a", "none", "null"):
            medidas = ""
        ws.cell(row=r, column=57, value=medidas)

        raw_weight = item.get('peso_g')
        grams = weight_grams(raw_weight)
        if raw_weight not in (None, '', '-') and grams is None:
            raise ValueError(f'Peso del producto inválido para {sku}: introduce gramos positivos')
        if grams is not None:
            if not all(weight_columns.values()):
                raise ValueError('La plantilla no contiene valor y unidad de peso del artículo')
            ws.cell(row=r, column=weight_columns['value'], value=grams)
            ws.cell(row=r, column=weight_columns['unit'], value='Gramos')

        # Medidas detalladas del producto (Grosor, Altura, Ancho y sus Unidades)
        dim_parts = parse_dimension_parts(medidas)
        if dim_parts:
            # Col DQ (121): Grosor del artículo desde la parte delantera hasta la trasera
            ws.cell(row=r, column=121, value=dim_parts["depth"])
            # Col DR (122): Unidad de grosor del artículo
            ws.cell(row=r, column=122, value=dim_parts["unit"])

            # Col DS (123): Altura del artículo desde la base hasta la parte superior
            ws.cell(row=r, column=123, value=dim_parts["height"])
            # Col DT (124): Unidad de altura del artículo
            ws.cell(row=r, column=124, value=dim_parts["unit"])

            # Col DU (125): Ancho del artículo de lado a lado
            ws.cell(row=r, column=125, value=dim_parts["width"])
            # Col DV (126): Unidad del ancho del artículo
            ws.cell(row=r, column=126, value=dim_parts["unit"])

        # Col EF (136): Estado del producto
        ws.cell(row=r, column=136, value="Nuevo")

        # Col FH (164): Tiempo de tramitación
        ws.cell(row=r, column=164, value=2)

        # Col FO (171): Fecha de finalización de la venta
        ws.cell(row=r, column=171, value="2035-12-31")

        # Col FP (172): Fecha de comienzo de la venta
        ws.cell(row=r, column=172, value="2026-08-01")

        # Col GJ (192): Plantilla de envío
        ws.cell(row=r, column=192, value="ENVIO 4")

        # Col JU (281): País de origen
        ws.cell(row=r, column=281, value="España")

        # Col JV (282): ¿Se necesitan baterías?
        ws.cell(row=r, column=282, value="No")

        # Col MH (346): Certificación de seguridad GPSR
        ws.cell(row=r, column=346, value="Sí")

        # Col MI (347): Dirección electrónica o de correo electrónico del fabricante
        ws.cell(row=r, column=347, value="vidalregals@gmail.com")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
